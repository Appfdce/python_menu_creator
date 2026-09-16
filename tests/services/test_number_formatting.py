import pytest
from docx import Document
from app.schemas.estimate_total import EstimateTotalRequest
from app.services.estimate_docx_generator import EstimateDocxGenerator
from app.services.estimate_perday_docx_generator import EstimatePerDayDocxGenerator


payload = {
    "event_id": "TEST-NUMFMT",
    "client": {"name": "Test Client"},
    "client_representative": {"name": "Rep Name"},
    "event": {
        "name": "Number Format Test",
        "guests": 10,
        "date_formatted": "Tuesday, Oct 27 2026",
        "end_date_formatted": "Friday, Oct 30 2026",
    },
    "meals": [
        {
            "date_header": "October, Tuesday 27 2026",
            "category_name": "Lunch",
            "category_precio_guest": "Lunch",
            "total_category_precio": "100",
            "total_category_precio_guest_por_dia": "100",
            "guest_count": 10,
            "provide_by_client": False,
            "subcategory_1_items": [{"name": "Eggs"}],
        }
    ],
    "labor_services": [
        {
            "date_header": "October, Tuesday 27 2026",
            "hours": "5,00",
            "name": "Bartender",
            "total": "50",
        }
    ],
    "extras_events": [],
    "financials": {
        "tax_rate": "8,875%",
        "tax_name": "Tax",
        "service_charge_rate": "10,00 %",
        "credit_card_percent": "0,00 %",
    },
}


GENERATORS = [EstimateDocxGenerator, EstimatePerDayDocxGenerator]


@pytest.mark.parametrize("generator_cls", GENERATORS)
def test_format_number(generator_cls):
    g = generator_cls()
    assert g._format_number("5,00") == "5"
    assert g._format_number("5.0") == "5"
    assert g._format_number(5.0) == "5"
    assert g._format_number("5") == "5"
    assert g._format_number("5,5") == "5.5"
    assert g._format_number("5.5") == "5.5"
    assert g._format_number("8,875") == "8.875"
    assert g._format_number(None) == ""
    assert g._format_number("") == ""


@pytest.mark.parametrize("generator_cls", GENERATORS)
def test_format_rate(generator_cls):
    g = generator_cls()
    assert g._format_rate("10,00 %") == "10%"
    assert g._format_rate("18%") == "18%"
    assert g._format_rate("8,875%") == "8.875%"
    assert g._format_rate("0,000 %") == "0%"
    assert g._format_rate(None) == ""
    assert g._format_rate("") == ""


def _full_text(docx_bytes):
    doc = Document(docx_bytes)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


@pytest.mark.parametrize("generator_cls", GENERATORS)
def test_docx_drops_trailing_decimals(generator_cls):
    g = generator_cls()
    req = EstimateTotalRequest(**payload)
    text = _full_text(g.generate_docx(req))

    # Labor hours: integer, no decimals
    assert "5 hours of labor." in text
    assert "5,00 hours" not in text
    assert "5.00 hours" not in text

    # Rates: no trailing decimals, no space before the '%'
    assert "10% Service Charge" in text
    assert "8.875% Tax" in text
    assert "10,00 %" not in text
