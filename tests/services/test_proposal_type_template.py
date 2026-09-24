import pytest
from docx import Document
from app.schemas.estimate_total import EstimateTotalRequest
from app.services.estimate_docx_generator import EstimateDocxGenerator
from app.services.estimate_perday_docx_generator import EstimatePerDayDocxGenerator

# Base payload shared by both modalities. The proposal_type is injected per test.
base_payload = {
    "event_id": "TEST-PROPOSAL-TYPE",
    "client": {"name": "Test Client"},
    "client_representative": {"name": "Rep Name"},
    "event": {
        "name": "Proposal Type Test",
        "guests": 50,
        "date_formatted": "Monday, June 15 2026",
        "end_date_formatted": "Monday, June 15 2026",
    },
    "meals": [
        {
            "show_date_header": True,
            "date_header": "June, Monday 15 2026",
            "category_name": "BREAKFAST",
            "category_precio_guest": "BREAKFAST @ $30.00 per Guest",
            "total_category_precio": "1500",
            "total_category_precio_guest_por_dia": "1500",
            "guest_count": "50",
            "date_day_name": "Monday",
        }
    ],
    "labor_services": [],
    "extras_events": [],
    "financials": {
        "tax_rate": "0",
        "tax_name": "Tax",
        "service_charge_rate": "0",
        "credit_card_percent": "0",
    },
}

GENERATORS = {
    "estimate": EstimateDocxGenerator,
    "perday": EstimatePerDayDocxGenerator,
}

# Text that only exists in the last page of each variant.
ANNUAL_MARKERS = ["Governing Terms", "Master Caterer Agreement"]
STANDARD_MARKERS = ["Events Policies", "Deposit"]


def _all_text(docx_stream):
    doc = Document(docx_stream)
    parts = [p.text for p in doc.paragraphs]

    def collect_tables(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
                    if cell.tables:
                        collect_tables(cell.tables)

    collect_tables(doc.tables)
    return "\n".join(parts)


def _payload_with(proposal_type):
    payload = {**base_payload, "event": {**base_payload["event"]}}
    payload["event"]["proposal_type"] = proposal_type
    return payload


@pytest.mark.parametrize("generator_cls", GENERATORS.values(), ids=GENERATORS.keys())
def test_annual_proposal_type_uses_annual_template(generator_cls):
    req = EstimateTotalRequest(**_payload_with("Annual"))
    text = _all_text(generator_cls().generate_docx(req))

    for marker in ANNUAL_MARKERS:
        assert marker in text, f"Missing annual marker '{marker}'"
    for marker in STANDARD_MARKERS:
        assert marker not in text, f"Unexpected standard marker '{marker}'"


@pytest.mark.parametrize("generator_cls", GENERATORS.values(), ids=GENERATORS.keys())
def test_standard_proposal_type_uses_standard_template(generator_cls):
    req = EstimateTotalRequest(**_payload_with("Standard"))
    text = _all_text(generator_cls().generate_docx(req))

    for marker in STANDARD_MARKERS:
        assert marker in text, f"Missing standard marker '{marker}'"
    for marker in ANNUAL_MARKERS:
        assert marker not in text, f"Unexpected annual marker '{marker}'"


@pytest.mark.parametrize("generator_cls", GENERATORS.values(), ids=GENERATORS.keys())
def test_empty_proposal_type_falls_back_to_legacy_template(generator_cls):
    req = EstimateTotalRequest(**_payload_with(""))
    generator = generator_cls()
    # Empty type must resolve to the legacy template.
    assert generator._resolve_template_path(req.event.proposal_type).endswith("_template.docx")
    # And generation must still succeed.
    assert generator.generate_docx(req).getvalue()


@pytest.mark.parametrize("generator_cls", GENERATORS.values(), ids=GENERATORS.keys())
def test_missing_variant_template_falls_back_to_legacy(generator_cls, monkeypatch):
    import app.services.estimate_docx_generator as standard_module
    import app.services.estimate_perday_docx_generator as perday_module

    module = standard_module if generator_cls is EstimateDocxGenerator else perday_module
    monkeypatch.setitem(
        module.TEMPLATE_BY_PROPOSAL_TYPE, "annual", "/tmp/does_not_exist_annual.docx"
    )

    req = EstimateTotalRequest(**_payload_with("Annual"))
    generator = generator_cls()
    # Missing variant file must degrade to the legacy template, not crash.
    assert generator._resolve_template_path(req.event.proposal_type).endswith("_template.docx")
    assert generator.generate_docx(req).getvalue()


def test_proposal_type_normalization():
    assert EstimateTotalRequest(**_payload_with("Annual")).event.proposal_type == "annual"
    assert EstimateTotalRequest(**_payload_with("annual")).event.proposal_type == "annual"
    assert EstimateTotalRequest(**_payload_with("Anual")).event.proposal_type == "annual"
    assert EstimateTotalRequest(**_payload_with("Standard")).event.proposal_type == "standard"
    assert EstimateTotalRequest(**_payload_with("STANDARD")).event.proposal_type == "standard"
    assert EstimateTotalRequest(**_payload_with("Particular")).event.proposal_type == "standard"
    assert EstimateTotalRequest(**_payload_with("")).event.proposal_type == ""
    assert EstimateTotalRequest(**_payload_with("whatever")).event.proposal_type == ""
