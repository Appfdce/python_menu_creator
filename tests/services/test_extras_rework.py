import pytest
from docx import Document
from app.schemas.estimate_total import EstimateTotalRequest
from app.services.estimate_docx_generator import EstimateDocxGenerator
from app.services.estimate_perday_docx_generator import EstimatePerDayDocxGenerator

payload = {
    "event_id": "TEST-EXTRAS-REWORK",
    "client": {"name": "Test Client"},
    "client_representative": {"name": "Rep Name"},
    "event": {
        "name": "Extras Rework Event",
        "guests": 10,
        "date_formatted": "Tuesday, Oct 27 2026",
        "end_date_formatted": "Friday, Oct 30 2026",
    },
    "meals": [],
    "labor_services": [],
    "extras_events": [
        # Day 1 - rentals
        {"date_header": "October, Tuesday 27 2026", "rental": True, "name": "Tent", "price": "1000", "qty": "1"},
        {"date_header": "October, Tuesday 27 2026", "rental": True, "name": "Chairs", "price": "200", "qty": "3"},
        # Day 1 - sales
        {"date_header": "October, Tuesday 27 2026", "rental": False, "name": "Flowers", "price": "50", "qty": "2"},
        # Day 1 - provided by client (excluded from totals, shown as such)
        {"date_header": "October, Tuesday 27 2026", "rental": True, "name": "Stage", "price": "999", "qty": "1", "provide_by_client": True},
        # Day 2 - only sales
        {"date_header": "October, Wednesday 28 2026", "rental": False, "name": "Linens", "price": "75", "qty": "1"},
    ],
    "financials": {
        "tax_rate": "0",
        "tax_name": "Tax",
        "service_charge_rate": "0",
        "credit_card_percent": "0",
    },
}


def _gen_text(generator_cls):
    req = EstimateTotalRequest(**payload)
    docx_bytes = generator_cls().generate_docx(req)
    doc = Document(docx_bytes)
    return "\n".join(p.text for p in doc.paragraphs)


def _gen_paragraphs(generator_cls):
    req = EstimateTotalRequest(**payload)
    docx_bytes = generator_cls().generate_docx(req)
    doc = Document(docx_bytes)
    return [p.text for p in doc.paragraphs]


@pytest.mark.parametrize("generator_cls", [EstimateDocxGenerator, EstimatePerDayDocxGenerator])
def test_extras_grouped_by_date_and_type(generator_cls):
    text = _gen_text(generator_cls)

    # Both dates appear as headers
    assert "Tuesday, October 27th, 2026" in text
    assert "Wednesday, October 28th, 2026" in text

    # Each extra line shows name + individual price x qty
    assert "Tent\t$ 1,000.00" in text
    assert "Chairs\t$ 600.00" in text
    assert "Flowers\t$ 100.00" in text
    assert "Linens\t$ 75.00" in text

    # Provided-by-client shows the placeholder text, not a price
    assert "Stage\tProvide by the client" in text
    assert "999" not in text.split("Provide by the client")[0]


@pytest.mark.parametrize("generator_cls", [EstimateDocxGenerator, EstimatePerDayDocxGenerator])
def test_extras_financial_totals_from_price(generator_cls):
    text = _gen_text(generator_cls)

    # Rentals = Tent(1000*1) + Chairs(200*3) = 1600 (Stage excluded)
    assert "Extras Services (Rentals)" in text
    assert "$ 1,600.00" in text

    # Sales = Flowers(50*2) + Linens(75*1) = 175
    assert "Extras Services (Sales)" in text
    assert "$ 175.00" in text


@pytest.mark.parametrize("generator_cls", [EstimateDocxGenerator, EstimatePerDayDocxGenerator])
def test_extras_headers_by_date(generator_cls):
    paras = _gen_paragraphs(generator_cls)

    # Standalone section headers (exact match) are only the ones inside Extras Services.
    # Rentals header appears once (only on Oct 27, since Oct 28 has no rentals).
    # Sales header appears twice (Oct 27 and Oct 28, both have sales).
    assert paras.count("Rentals") == 1, f"Expected exactly 1 'Rentals' header, got {paras.count('Rentals')}"
    assert paras.count("Sales") == 2, f"Expected exactly 2 'Sales' headers, got {paras.count('Sales')}"


@pytest.mark.parametrize("generator_cls", [EstimateDocxGenerator, EstimatePerDayDocxGenerator])
def test_extras_qty_defaults_to_one(generator_cls):
    req = EstimateTotalRequest(**{
        "event_id": "TEST-QTY",
        "client": {"name": "C"},
        "client_representative": {"name": "R"},
        "event": {"name": "Qty Event", "guests": 5,
                  "date_formatted": "Tuesday, Oct 27 2026",
                  "end_date_formatted": "Friday, Oct 30 2026"},
        "meals": [],
        "labor_services": [],
        "extras_events": [
            {"date_header": "October, Tuesday 27 2026", "rental": False, "name": "NoQty", "price": "10"},
        ],
        "financials": {"tax_rate": "0", "tax_name": "Tax", "service_charge_rate": "0",
                       "credit_card_percent": "0"},
    })
    docx_bytes = generator_cls().generate_docx(req)
    doc = Document(docx_bytes)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "NoQty\t$ 10.00" in text
